import os
from flask import Flask, render_template, request, send_file
import pdfplumber
import docx
from werkzeug.utils import secure_filename
import google.generativeai as genai
from fpdf import FPDF 
import re
os.environ["GOOGLE_API_KEY"] = "---------------" 
genai.configure(api_key=os.environ["GOOGLE_API_KEY"])
model = genai.GenerativeModel("gemini-2.5-pro")
app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads/'
app.config['RESULTS_FOLDER'] = 'results/'
app.config['ALLOWED_EXTENSIONS'] = {'pdf', 'txt', 'docx'}
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']
def extract_text_from_file(file_path):
    ext = file_path.rsplit('.', 1)[1].lower()
    if ext == 'pdf':
        try:
            with pdfplumber.open(file_path) as pdf:
                text = ''.join([page.extract_text() or '' for page in pdf.pages])
            return text
        except Exception as e:
            print(f"PDF extraction error: {e}")
            return None
    elif ext == 'docx':
        try:
            doc = docx.Document(file_path)
            return ' '.join([para.text for para in doc.paragraphs])
        except Exception as e:
            print(f"DOCX extraction error: {e}")
            return None
    elif ext == 'txt':
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"TXT reading error: {e}")
            return None
    return None
def generate_questions(input_text, num_mcqs=5, num_short=5, num_long=3):
    """
    Generates multiple MCQs, short answers, and long answers using AI.
    """
    prompt = f"""
You are an AI assistant. Generate a professional exam paper from the text below:

Text: {input_text}

Requirements:
1. Generate exactly {num_mcqs} MCQs, each with 4 options (A-D) and correct answer.
2. Generate exactly {num_short} short answer questions (2 marks each).
3. Generate exactly {num_long} long answer questions.

Output format:
- Each MCQ must start with "## MCQ"
- Each short answer must start with "## SHORT"
- Each long answer must start with "## LONG"

Example:

## MCQ
Question: <question text>
A) option A
B) option B
C) option C
D) option D
Correct Answer: <A/B/C/D>

## SHORT
Question: <short answer question text>

## LONG
Question: <long answer question text>

Repeat blocks to match the number of questions requested.
"""
    try:
        response = model.generate_content(prompt).text.strip()
        return response
    except Exception as e:
        print("Error generating questions:", e)
        return ""

def parse_questions(text):
    """
    Parses AI output text into structured Python dicts.
    """
    mcqs = []
    short_ans = []
    long_ans = []
    mcq_blocks = re.findall(r"## MCQ(.*?)(?=##|$)", text, re.DOTALL)
    short_blocks = re.findall(r"## SHORT(.*?)(?=##|$)", text, re.DOTALL)
    long_blocks = re.findall(r"## LONG(.*?)(?=##|$)", text, re.DOTALL)
    for block in mcq_blocks:
        try:
            q_match = re.search(r"Question:\s*(.*?)\s*[A-D]\)", block, re.DOTALL | re.IGNORECASE)
            q_text = q_match.group(1).strip() if q_match else None
            opt_A_match = re.search(r"A\)(.*?)\s*B\)", block, re.DOTALL | re.IGNORECASE)
            opt_B_match = re.search(r"B\)(.*?)\s*C\)", block, re.DOTALL | re.IGNORECASE)
            opt_C_match = re.search(r"C\)(.*?)\s*D\)", block, re.DOTALL | re.IGNORECASE)
            opt_D_match = re.search(r"D\)(.*?)\s*Correct Answer:", block, re.DOTALL | re.IGNORECASE)
            answer_match = re.search(r"Correct Answer:\s*(.*)", block, re.DOTALL | re.IGNORECASE)
            answer = answer_match.group(1).strip() if answer_match else "N/A"
            if q_text and opt_A_match and opt_B_match and opt_C_match and opt_D_match:
                options = {
                    "A": opt_A_match.group(1).strip(),
                    "B": opt_B_match.group(1).strip(),
                    "C": opt_C_match.group(1).strip(),
                    "D": opt_D_match.group(1).strip(),
                }
                mcqs.append({"question": q_text, "options": options, "answer": answer})
            else:
                print(f"Skipping malformed MCQ block (Missing components): {block[:100]}...")
        except Exception as e:
            print(f"Error parsing MCQ block: {e}. Block: {block[:100]}...")
            continue
    for block in short_blocks:
        try:
            q_text = re.search(r"Question:(.*)", block, re.DOTALL).group(1).strip()
            short_ans.append(q_text)
        except:
            continue
    for block in long_blocks:
        try:
            q_text = re.search(r"Question:(.*)", block, re.DOTALL).group(1).strip()
            long_ans.append(q_text)
        except:
            continue
    return {"mcqs": mcqs, "short": short_ans, "long": long_ans}
def save_text_file(content, filename):
    path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)
    return path
def create_pdf_file(structured_questions, filename, exam_info=None):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_left_margin(15)
    pdf.set_right_margin(15)
    if exam_info:
        pdf.set_font("Arial", "B", 16)
        pdf.cell(0, 10, f"Date: {exam_info.get('date', '')}", ln=0)
        pdf.cell(0, 10, f"Time: {exam_info.get('time', '')}", ln=1, align="R")
        pdf.set_font("Arial", "B", 16)
        pdf.cell(0, 10, f"Exam: {exam_info.get('title', '')}", ln=0)
        pdf.cell(0, 10, f"Total Marks: {exam_info.get('total_marks', '')}", ln=1, align="R")
        pdf.ln(10)
    pdf.set_font("Arial", "B", 18)
    pdf.cell(0, 10, "Question Paper", ln=True, align="C")
    pdf.ln(5)
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Section A", ln=True, align="C")
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "**Multiple Choice Questions (MCQs)**", ln=True, align="C")
    pdf.set_font("Arial", "", 12)
    pdf.ln(5)
    if not structured_questions["mcqs"]:
        pdf.cell(0, 10, "No MCQs generated.", ln=True)
        pdf.ln(5)
    for idx, mcq in enumerate(structured_questions["mcqs"], 1):
        pdf.multi_cell(0, 8, f"{idx}. {mcq['question']}")
        pdf.ln(1)  
        for opt in ["A", "B", "C", "D"]:
            if opt in mcq['options']:
                pdf.multi_cell(0, 8, f"{opt}) {mcq['options'][opt]}")
        pdf.ln(2)
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Section B", ln=True, align="C")
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "**Short Answer Questions - 2 Marks**", ln=True, align="C")
    pdf.set_font("Arial", "", 12)
    pdf.ln(5)
    for idx, q in enumerate(structured_questions["short"], 1):
        pdf.multi_cell(0, 8, f"{idx}. {q}")
        pdf.ln(2)
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Section C", ln=True, align="C")
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "**Long Answer Questions**", ln=True, align="C")
    pdf.set_font("Arial", "", 12)
    pdf.ln(5)
    total_long = len(structured_questions["long"])
    if total_long > 0:
        pdf.multi_cell(0, 8, f"How many questions: {total_long}")
        pdf.multi_cell(0, 8, f"How many need to answer: {total_long - 1 if total_long > 1 else 1}")
        pdf.ln(5)
    for idx, q in enumerate(structured_questions["long"], 1):
        pdf.multi_cell(0, 8, f"{idx}. {q}")
        pdf.ln(2)
    path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    pdf.output(path)
    return path
@app.route('/')
def index():
    return render_template("index.html")
@app.route('/generate', methods=['POST'])
def generate_exam():
    file = request.files.get('file')
    text = request.form.get('input_text') or ""
    file_path = None
    filename = ""
    if file and file.filename and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)
        text_from_file = extract_text_from_file(file_path)
        if text_from_file:
            text = text_from_file
    if not text:
        return "Failed to extract text or no text provided."
    try:
        num_mcqs = int(request.form.get("section_a_count", 0))
        num_short = int(request.form.get("section_b_count", 0))
        num_long = int(request.form.get("section_c_count", 0))
    except ValueError:
        return "Invalid input for number of questions. Please enter integers."
    exam_info = {
        "title": request.form.get("exam", "Exam Paper"),  
        "date": request.form.get("date", ""),  
        "time": request.form.get("time", ""), 
        "course_code": request.form.get("course_code", ""),  
        "course_name": request.form.get("course_name", ""),
        "total_marks": request.form.get("total_marks", ""),
    }
    ai_text = generate_questions(text, num_mcqs, num_short, num_long)
    print("AI Output:", ai_text)
    structured_questions = parse_questions(ai_text)
    base_filename = filename.rsplit('.', 1)[0] if file_path else "generated_exam"
    txt_filename = f"exam_{base_filename}.txt"
    pdf_filename = f"exam_{base_filename}.pdf"
    save_text_file(ai_text, txt_filename)
    create_pdf_file(structured_questions, pdf_filename, exam_info)
    return render_template("results.html",
                           mcqs=structured_questions["mcqs"],
                           short_answers=structured_questions["short"],
                           long_answers=structured_questions["long"],
                           date=exam_info.get("date", ""),
                           exam=exam_info.get("title", ""),  
                           time=exam_info.get("time", ""),
                           course_code=exam_info.get("course_code", ""),  
                           course_name=exam_info.get("course_name", ""),  
                           total_marks=exam_info.get("total_marks", ""),
                           txt_filename=txt_filename,
                           pdf_filename=pdf_filename)
@app.route('/download/<filename>')
def download_file(filename):
    path = os.path.join(app.config['RESULTS_FOLDER'], filename)
    return send_file(path, as_attachment=True)
if __name__ == "__main__":
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    os.makedirs(app.config['RESULTS_FOLDER'], exist_ok=True)
    app.run(debug=True)